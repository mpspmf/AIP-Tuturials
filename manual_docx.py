# -*- coding: utf-8 -*-
"""Gera o manual editável em .docx (conteúdo alinhado com o PDF)."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "materiais"
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAMS = RGBColor(0x46, 0x4E, 0xB8)
MUTED = RGBColor(0x5C, 0x63, 0x70)


def _set_run_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def add_rich_paragraph(doc, text, style=None, size=11, align=None, space_after=6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            _set_run_font(run, size=size, bold=True)
        else:
            run = p.add_run(part)
            _set_run_font(run, size=size)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, size=sizes.get(level, 12), bold=True, color=NAVY if level == 1 else TEAMS)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        parts = re.split(r"(\*\*.*?\*\*)", item)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                _set_run_font(run, size=size, bold=True)
            else:
                run = p.add_run(part)
                _set_run_font(run, size=size)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                _set_run_font(run, size=10, bold=True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    _set_run_font(run, size=10)
    doc.add_paragraph()


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            _set_run_font(run, size=10.5, bold=True, color=NAVY)
        else:
            run = p.add_run(part)
            _set_run_font(run, size=10.5, color=NAVY)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    _set_run_font(run, size=9, italic=True, color=MUTED)
    p.paragraph_format.space_after = Pt(10)


def build_docx():
    path = OUT / "Manual_Microsoft_Teams_AIP.docx"
    try:
        # Detect lock early by attempting open for append
        with open(path, "a"):
            pass
    except PermissionError:
        path = OUT / "Manual_Microsoft_Teams_AIP_editavel.docx"
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Capa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COMUNICAÇÃO INTERNA")
    _set_run_font(r, size=12, bold=True, color=TEAMS)
    doc.add_paragraph()
    for line in [
        "Microsoft Teams e correio eletrónico:",
        "modelo de comunicação interna",
        "por área operacional",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        _set_run_font(r, size=20, bold=True, color=NAVY)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Manual de utilização para colaboradores e diretores de área operacional")
    _set_run_font(r, size=12, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AIP  ·  Versão 1.0  ·  2026")
    _set_run_font(r, size=11, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "Índice")
    for item in [
        "1. Finalidade deste manual",
        "2. Situação atual e modelo pretendido",
        "3. Critérios de utilização: Teams e correio eletrónico",
        "4. Organização: um Team por área operacional (AO)",
        "5. Canais recomendados e nomenclatura",
        "6. Guia prático de utilização",
        "7. Envio simultâneo para e-mail e canal do Teams",
        "8. Vantagens do Teams na comunicação operacional",
        "9. Orientações para a adoção progressiva",
        "10. Boas práticas de comunicação",
        "11. Notificações, telemóvel e disponibilidade",
        "12. Plano de implementação (90 dias)",
        "13. Perguntas frequentes",
        "14. Glossário",
    ]:
        add_rich_paragraph(doc, item, size=11, space_after=4)
    doc.add_page_break()

    # 1
    add_heading(doc, "1. Finalidade deste manual")
    add_rich_paragraph(
        doc,
        "O presente documento define o modelo de utilização do **Microsoft Teams** "
        "como instrumento principal de comunicação interna operacional na AIP, "
        "reservando o **correio eletrónico** para assuntos formais, comunicação externa "
        "e documentação com valor de arquivo.",
    )
    add_rich_paragraph(
        doc,
        "Não se pretende a eliminação do e-mail. Pretende-se a seleção adequada do canal "
        "em função da natureza de cada comunicação, com vista a reduzir tempos de resposta "
        "e a evitar a dispersão de informação em cadeias de mensagens.",
    )
    add_heading(doc, "Objetivos", 2)
    add_bullets(
        doc,
        [
            "Agilizar a comunicação **no âmbito de cada área operacional (AO)**.",
            "Atribuir a cada diretor de AO um espaço próprio para organização dos canais de trabalho.",
            "Manter o e-mail para comunicação com clientes, contratos, decisões formais e arquivo.",
            "Apoiar a transição de colaboradores habituados exclusivamente ao Outlook.",
        ],
    )
    add_callout(
        doc,
        "**Leitura recomendada.** As secções 2 a 5 apresentam o modelo organizacional. "
        "A secção 6 descreve os procedimentos de utilização. As secções 9 e 12 destinam-se "
        "prioritariamente a diretores e responsáveis pela implementação.",
    )

    # 2
    add_heading(doc, "2. Situação atual e modelo pretendido")
    add_heading(doc, "Situação atual", 2)
    add_rich_paragraph(
        doc,
        "O Microsoft Teams **já se encontra disponível** na organização, porém organizado "
        "num único espaço transversal: **um Team com dois canais**, abrangendo todos os "
        "colaboradores. Esta estrutura funciona como quadro de avisos gerais — adequada "
        "a comunicados pontuais, insuficiente para o trabalho quotidiano de cada AO. "
        "Consequentemente, a comunicação operacional permanece no e-mail, com os "
        "inconvenientes associados: demora, fragmentação e dificuldade de acompanhamento.",
    )
    add_heading(doc, "Modelo pretendido", 2)
    add_rich_paragraph(
        doc,
        "Transitar de «um Team institucional» para **um Team por AO**. "
        "O diretor da AO cria e gere os canais necessários (projetos, turnos, qualidade, "
        "esclarecimentos, entre outros). O Team institucional mantém-se exclusivamente "
        "para comunicados de interesse transversal.",
    )
    add_table(
        doc,
        ["Situação atual", "Modelo pretendido"],
        [
            ["1 Team para toda a organização", "1 Team por AO, sob gestão do diretor"],
            ["2 canais genéricos", "Canais por tema, projeto ou processo"],
            ["Comunicação operacional sobretudo por e-mail", "Comunicação operacional no Teams; e-mail para o formal"],
            ["Contexto disperso em cadeias e anexos", "Histórico centralizado no canal adequado"],
            ["Integração de novos colaboradores sem contexto", "Acesso ao histórico da AO no respetivo Team"],
        ],
    )
    add_caption(doc, "Quadro 1 — Comparação entre o modelo atual e o modelo pretendido.")

    # 3
    add_heading(doc, "3. Critérios de utilização: Teams e correio eletrónico")
    add_rich_paragraph(
        doc,
        "**Critério geral:** quando a comunicação for interna, operacional e exigir "
        "resposta célere, utiliza-se o **Teams**. Quando exigir formalidade, prova "
        "documental ou destinatários externos, utiliza-se o **correio eletrónico**.",
    )
    add_table(
        doc,
        ["Situação", "Canal", "Justificação"],
        [
            ["Pedido de esclarecimento à equipa da AO", "Teams — canal ou conversa", "Resposta no contexto; visibilidade dos intervenientes relevantes"],
            ["Acompanhamento de projeto interno", "Teams — canal do projeto", "Histórico único; ficheiros partilhados"],
            ["Esclarecimento pontual entre colegas", "Teams — conversa (chat)", "Substitui o e-mail de carácter operacional"],
            ["Reunião interna ou ponto de situação", "Teams — reunião", "Convite, vídeo, gravação e notas na mesma plataforma"],
            ["Edição colaborativa de documentos", "Teams + SharePoint da equipa", "Uma versão de trabalho; eliminação de anexos sucessivos"],
            [
                "Comunicado institucional ou aviso crítico da AO",
                "E-mail e canal (envio simultâneo)",
                "Garantia de alcance a utilizadores de Outlook e de Teams (secção 7)",
            ],
            ["Cliente, fornecedor ou entidade externa", "E-mail", "Canal profissional adequado à comunicação externa"],
            ["Contrato, faturação, decisão formal, auditoria", "E-mail (e repositório oficial, se existir)", "Rasto formal e valor de arquivo"],
            ["Assuntos de RH ou dados pessoais sensíveis", "E-mail / processo RH definido", "Confidencialidade e formalismo"],
            ["Confirmação escrita de decisão tomada no Teams", "E-mail breve de síntese", "O Teams permite a discussão; o e-mail formaliza quando necessário"],
        ],
    )
    add_caption(doc, "Quadro 2 — Matriz de seleção do canal de comunicação.")
    add_callout(
        doc,
        "**Orientação prática.** Tratando-se de um assunto interno com expectativa "
        "de resposta no próprio dia ou no dia seguinte, deve privilegiar-se o canal "
        "da AO ou uma conversa no Teams. Caso, no termo da discussão, seja necessária "
        "comprovação formal, emite-se um e-mail breve que sintetize o acordo alcançado.",
    )

    # 4
    add_heading(doc, "4. Organização: um Team por área operacional (AO)")
    add_rich_paragraph(
        doc,
        "Cada AO dispõe do seu próprio Team no Microsoft Teams. "
        "Trata-se do espaço de trabalho da área — distinto de uma lista de distribuição de e-mail.",
    )
    add_heading(doc, "Responsabilidades", 2)
    add_bullets(
        doc,
        [
            "**Diretor da AO:** responsável pelo Team. Cria e arquiva canais, define a composição, "
            "dá o exemplo (responde no Teams) e designa um embaixador.",
            "**Embaixador da AO (recomendado):** colaborador que apoia a equipa nos primeiros 30 dias e recolhe dúvidas.",
            "**Colaboradores da AO:** utilizam os canais para a atividade corrente; configuram notificações; "
            "não criam canais sem autorização.",
            "**IT / administração Microsoft 365:** assegura a criação inicial dos Teams (quando necessário), "
            "licenças e permissões. A gestão corrente compete à AO.",
        ],
    )
    add_heading(doc, "Team institucional", 2)
    add_rich_paragraph(
        doc,
        "O Team atual (toda a organização, dois canais) **não deve ser eliminado**. "
        "O seu âmbito passa a ser restrito: avisos institucionais, campanhas internas "
        "e informação de interesse transversal. Deixa de constituir o espaço para "
        "questões operacionais ou planeamento específico de uma AO.",
    )
    add_rich_paragraph(
        doc,
        "Caso os dois canais atuais tenham designações genéricas, recomenda-se: "
        "um canal **Avisos** (publicação pela direção / comunicação interna) e um canal "
        "**Geral** (comunicação pontual de caráter institucional). A restante atividade "
        "migra para o Team da respetiva AO.",
    )

    # 5
    add_heading(doc, "5. Canais recomendados e nomenclatura")
    add_rich_paragraph(
        doc,
        "Recomenda-se a utilização de designações curtas, consistentes entre AOs, "
        "sempre que tal se justifique. Esta uniformidade facilita a integração de "
        "colaboradores que mudem de área.",
    )
    add_table(
        doc,
        ["Canal", "Finalidade", "Quem publica"],
        [
            ["Geral", "Canal predefinido da AO. Coordenação operacional quotidiana.", "Todos os membros da AO"],
            ["Avisos", "Comunicados do diretor. Informação unidirecional e clara.", "Diretor (e quem for por si autorizado)"],
            ["Dúvidas", "Pedidos de esclarecimento. Substitui o e-mail operacional breve.", "Todos os membros da AO"],
            ["Projeto-[nome]", "Canal por iniciativa com calendário definido.", "Equipa do projeto"],
            ["Turnos / Operação", "Passagem de informação, ocorrências e prioridades do serviço.", "Operação"],
            ["Qualidade / Segurança", "Não conformidades, lembretes e boas práticas.", "Operação e qualidade"],
        ],
    )
    add_caption(doc, "Quadro 3 — Estrutura inicial de canais (adaptável pelo diretor da AO).")
    add_heading(doc, "Critérios para criação de canais", 2)
    add_bullets(
        doc,
        [
            "Criar um canal apenas quando o tema for **recorrente** ou constituir um projeto autónomo.",
            "Assuntos de curta duração devem ser tratados em conversa (thread) no canal Geral.",
            "Arquivar (sem eliminação imediata) canais de projetos concluídos.",
            "Recorrer a canais privados apenas quando exista confidencialidade efetiva.",
        ],
    )

    # 6
    add_heading(doc, "6. Guia prático de utilização")
    add_rich_paragraph(
        doc,
        "O Teams encontra-se já disponível. Para a maioria dos colaboradores, "
        "não se trata de instalar uma nova ferramenta, mas de **passar a utilizar "
        "o espaço adequado** a cada tipo de comunicação.",
    )
    add_heading(doc, "6.1 Acesso ao Team da AO", 2)
    add_bullets(
        doc,
        [
            "Abrir o **Microsoft Teams** no computador (ou a aplicação no telemóvel).",
            "Autenticar-se com a **conta profissional** da AIP (a mesma do Outlook).",
            "No menu esquerdo, selecionar **Equipas** (Teams).",
            "Localizar o Team correspondente à respetiva AO. Caso não esteja visível, "
            "solicitar a inclusão ao diretor ou ao IT — evitando, entretanto, "
            "recorrer ao e-mail operacional por omissão.",
            "Abrir o canal **Geral** e consultar as publicações recentes. "
            "Este constitui o espaço principal de coordenação da área.",
        ],
    )
    add_heading(doc, "6.2 Publicação no canal", 2)
    add_bullets(
        doc,
        [
            "Utilizar a caixa de mensagem situada na parte inferior do canal.",
            "Redigir de forma objetiva: contexto, pedido e prazo, quando aplicável.",
            "Utilizar **@nome** para identificar o destinatário responsável. "
            "Utilizar **@canal** apenas quando a informação for indispensável a todos os membros.",
            "Para dar continuidade ao mesmo assunto, utilizar **Responder** à mensagem original. "
            "Evitar iniciar uma nova publicação sobre o mesmo tema.",
            "Carregar o ficheiro no canal, em substituição do envio por anexo de e-mail, "
            "de modo a manter uma única versão no local adequado.",
        ],
    )
    add_heading(doc, "6.3 Conversas (chat)", 2)
    add_rich_paragraph(
        doc,
        "A conversa (chat) deve ser utilizada para comunicações breves entre uma ou "
        "duas pessoas — esclarecimentos pontuais que anteriormente ocorriam por e-mail. "
        "Quando o assunto passar a interessar à AO, **deve ser remetido ao canal** "
        "com um resumo objetivo. O chat privado não substitui o histórico da equipa.",
    )
    add_heading(doc, "6.4 Reuniões", 2)
    add_bullets(
        doc,
        [
            "No canal ou no calendário do Teams, selecionar **Reunir** / agendar.",
            "É admissível criar o convite a partir do Outlook; a reunião realiza-se no Teams.",
            "Partilhar o ecrã para apresentação de documentos, em alternativa ao envio prévio de anexos.",
            "Decisões relevantes devem ser registadas no canal após a reunião.",
        ],
    )
    add_heading(doc, "6.5 Ficheiros", 2)
    add_rich_paragraph(
        doc,
        "Cada Team dispõe de uma biblioteca de ficheiros (SharePoint). "
        "Acesso: Teams → canal → separador **Ficheiros**. "
        "A edição em linha evita a circulação de múltiplas versões por anexo.",
    )
    add_heading(doc, "6.6 Pesquisa", 2)
    add_rich_paragraph(
        doc,
        "A barra de pesquisa do Teams localiza mensagens, ficheiros e pessoas. "
        "Antes de solicitar informação já disponível, recomenda-se uma pesquisa prévia. "
        "Caso a informação não seja encontrada, o pedido deve ser colocado no canal "
        "**Dúvidas**, de modo a ficar disponível para futuros colaboradores.",
    )
    add_heading(doc, "6.7 Ações iniciais do diretor da AO", 2)
    add_bullets(
        doc,
        [
            "Confirmar junto do IT a existência do Team da AO e a correta composição de membros.",
            "Criar os canais iniciais (Geral, Avisos, Dúvidas e os canais específicos da operação).",
            "Publicar no canal Geral o comunicado de arranque, com a data de início e os critérios "
            "de utilização (Teams para o operacional; e-mail para o formal e externo).",
            "Fixar essa publicação ou o presente manual no canal Avisos.",
            "Obter o **endereço de e-mail do canal Avisos** (secção 7) e registá-lo nos contactos do Outlook.",
            "Designar o embaixador e indicar o contacto de apoio durante a fase inicial.",
        ],
    )

    # 7
    add_heading(doc, "7. Envio simultâneo para e-mail e canal do Teams")
    add_rich_paragraph(
        doc,
        "Durante a transição — e sempre que um comunicado revista importância suficiente "
        "para não depender exclusivamente do acesso ao Teams — é possível **redigir uma única "
        "mensagem e assegurá-la em ambos os canais**: a caixa de entrada e o canal. "
        "Esta funcionalidade é particularmente útil para colaboradores que consultam "
        "predominantemente o Outlook.",
    )
    add_callout(
        doc,
        "**Utilização criteriosa.** A duplicação sistemática de todas as comunicações "
        "anula o benefício da mudança. O envio simultâneo deve reservar-se a avisos, "
        "prazos, alterações de turno, matérias de segurança e comunicados da direção. "
        "A comunicação operacional corrente deve decorrer exclusivamente no canal.",
    )
    add_heading(doc, "7.1 Método A — E-mail no Outlook com publicação automática no canal", 2)
    add_rich_paragraph(
        doc,
        "Cada canal do Teams dispõe de um **endereço de e-mail próprio**. "
        "Ao incluí-lo em **Para** ou em **Cc** no Outlook, os destinatários recebem "
        "o e-mail e a mesma mensagem é publicada no canal.",
    )
    add_bullets(
        doc,
        [
            "No Teams, abrir o canal (por exemplo **Avisos**). Selecionar as **opções** (…) junto ao nome do canal.",
            "Escolher **Obter endereço de e-mail** (Get email address).",
            "Copiar o endereço e, se conveniente, **guardá-lo como contacto** no Outlook com designação clara.",
            "No Outlook, redigir o e-mail. Em **Cc** (recomendado), incluir o endereço do canal, "
            "além dos destinatários habituais.",
            "Enviar. A mensagem é publicada no canal (assunto, corpo e anexos, dentro dos limites da Microsoft).",
        ],
    )
    add_rich_paragraph(
        doc,
        "O diretor (ou o IT) pode restringir quem está autorizado a enviar e-mail para "
        "determinado canal — por exemplo, apenas membros da AO. Recomenda-se essa "
        "restrição nos canais de avisos, por motivos de segurança.",
    )
    add_heading(doc, "7.2 Método B — Partilhar no Teams um e-mail existente", 2)
    add_rich_paragraph(
        doc,
        "Quando a mensagem **já consta** do Outlook (recebida ou enviada) e se pretende "
        "disponibilizá-la à AO no canal:",
    )
    add_bullets(
        doc,
        [
            "Abrir o e-mail no Outlook (aplicação de ambiente de trabalho ou na web).",
            "Na faixa de opções, selecionar **Partilhar no Teams** (Share to Teams). "
            "Caso a opção não esteja disponível, o IT poderá ter de a ativar na organização.",
            "Selecionar o Team e o **canal** da AO (e não o chat, quando o objetivo for o histórico da equipa).",
            "Acrescentar uma nota de contexto e concluir a partilha.",
        ],
    )
    add_heading(doc, "7.3 Método C — Publicação no canal com aviso por e-mail", 2)
    add_rich_paragraph(
        doc,
        "Percurso inverso: publica-se primeiro no canal e, subsequentemente, remete-se "
        "um e-mail breve com a ligação da conversa a quem ainda não consulta o Teams. "
        "No Teams: **Mais opções (…) → Copiar ligação**. Quem aceder à ligação "
        "abre o ponto exacto do canal (desde que disponha de acesso ao Team).",
    )
    add_heading(doc, "7.4 Limitações relevantes", 2)
    add_bullets(
        doc,
        [
            "**Não se trata de uma conversa sincronizada.** As respostas no Outlook "
            "permanecem no fio de e-mail; as do canal permanecem no Teams. "
            "Deve definir-se à partida o canal de continuação (preferencialmente o Teams).",
            "Existem **limites de dimensão e de anexos** no envio de e-mail para canal. "
            "Ficheiros de grande dimensão devem ser colocados na biblioteca do Team.",
            "Não devem ser enviados para o endereço do canal **dados pessoais sensíveis, "
            "assuntos de RH ou informação de clientes** cuja visibilidade não seja adequada "
            "a todos os membros do canal.",
            "O endereço do canal não deve ser divulgado fora da organização.",
        ],
    )
    add_table(
        doc,
        ["Situação", "Método", "Observações"],
        [
            ["Comunicado da AO de alcance obrigatório", "A — e-mail com o canal em Cc", "Um envio; dois destinos"],
            ["E-mail externo que a AO deve acompanhar", "B — Partilhar no Teams", "Não incluir o interlocutor externo no canal"],
            ["Publicação já efetuada; necessidade de alertar um colaborador", "C — e-mail com a ligação", "Evita a reescrita do conteúdo"],
            ["Esclarecimento operacional corrente", "Apenas Teams", "Sem duplicação no correio"],
        ],
    )
    add_caption(doc, "Quadro 4 — Critérios para envio simultâneo ou exclusivo.")

    # 8
    add_heading(doc, "8. Vantagens do Teams na comunicação operacional")
    add_rich_paragraph(
        doc,
        "O correio eletrónico foi concebido para correspondência formal. "
        "A atividade de uma AO caracteriza-se por comunicações paralelas, documentos "
        "em evolução e decisões céleres. O Teams responde a esse ritmo. Em particular:",
    )
    add_bullets(
        doc,
        [
            "**Celeridade:** a informação é disponibilizada onde a equipa já trabalha.",
            "**Contexto:** no canal do projeto, a conversa não se mistura com correspondência alheia ao tema.",
            "**Redução de «Responder a todos»:** apenas os membros do canal são abrangidos.",
            "**Versão única do ficheiro:** edição colaborativa; eliminação de anexos sucessivos.",
            "**Reuniões integradas** no mesmo ambiente em que se desenvolve o trabalho.",
            "**Indicador de presença:** possibilidade de verificar disponibilidade antes de contactar.",
            "**@menções:** identificação precisa do responsável pela ação.",
            "**Pesquisa centralizada** de mensagens e ficheiros da AO.",
            "**Memória organizacional:** novos colaboradores acedem ao histórico da área.",
            "**Acesso móvel:** a mesma plataforma no telemóvel, com notificações configuráveis (secção 11).",
            "**Integrações:** Planner, listas e aprovações poderão ser adotadas numa fase posterior.",
        ],
    )
    add_rich_paragraph(
        doc,
        "O e-mail permanece o instrumento adequado à comunicação externa e formal. "
        "A alteração de modelo não desvaloriza o correio eletrónico; "
        "atribui a cada ferramenta a finalidade para que foi concebida.",
    )

    # 9
    add_heading(doc, "9. Orientações para a adoção progressiva")
    add_rich_paragraph(
        doc,
        "A familiaridade com o Outlook constitui competência acumulada, não resistência "
        "injustificada. A adoção falha quando se solicita a mudança de ferramenta sem "
        "alterar as regras de utilização. Se a direção continuar a responder exclusivamente "
        "por e-mail, o Teams não será adotado.",
    )
    add_heading(doc, "Medidas recomendadas", 2)
    add_bullets(
        doc,
        [
            "**Exemplo da liderança.** Questões recebidas no Teams devem ser respondidas no Teams. "
            "Assuntos operacionais recebidos por e-mail podem ser remetidos ao canal adequado, "
            "com indicação clara. Nos primeiros 30 dias, comunicados relevantes podem ser "
            "enviados **em simultâneo** para o e-mail e para o canal (secção 7).",
            "**Critérios em formato resumido.** O quadro 2 deste manual deve ser disponibilizado "
            "no canal Avisos.",
            "**Resultados iniciais (primeiros 15 dias):** transferir comunicados internos para Avisos; "
            "esclarecimentos para Dúvidas; reunião semanal da AO no Teams.",
            "**Formação prática:** sessão de 20 minutos por AO, com demonstração das ações das "
            "secções 6 e 7. Gravar e disponibilizar no canal.",
            "**Embaixador:** colaborador com disponibilidade para apoio, não necessariamente com perfil técnico avançado.",
            "**Gestão de notificações:** configurar alertas essenciais. O excesso de notificações "
            "induz o regresso ao e-mail.",
            "**Transição respeitosa:** reconhecer o valor do e-mail formal e valorizar a utilização "
            "correta do canal, não o volume de mensagens.",
            "**Projeto-piloto (recomendado):** uma AO durante 3 a 4 semanas, ajuste da estrutura "
            "e posterior replicação. O Team institucional mantém-se para avisos transversais.",
        ],
    )
    add_heading(doc, "Formulação recomendada (diretores e embaixadores)", 2)
    add_bullets(
        doc,
        [
            "«Tratando-se de assunto operacional da AO, solicita-se a publicação no canal, "
            "para conhecimento da equipa.»",
            "«Para comunicação com o cliente, mantém-se o e-mail. Para comunicação interna, utiliza-se o Teams.»",
            "«Não se exige resposta imediata fora do expediente. Pretende-se que o e-mail interno "
            "fique reservado a assuntos formais.»",
            "«Caso a notificação não tenha sido visualizada, o histórico permanece disponível no canal.»",
        ],
    )
    add_callout(
        doc,
        "**Práticas a evitar.** Orientações genéricas sem critérios concretos; utilização de "
        "grupos de WhatsApp para operação; canais em que apenas a direção publica. "
        "O Teams só produz benefícios quando a comunicação é bidirecional.",
    )

    # 10
    add_heading(doc, "10. Boas práticas de comunicação")
    add_bullets(
        doc,
        [
            "Indicar claramente o assunto na primeira linha da mensagem.",
            "Concentrar um pedido de ação por mensagem, sempre que possível.",
            "Responder na conversa (thread) correspondente.",
            "Não utilizar @canal para cumprimentos ou partilhas sem relevância operacional.",
            "Assuntos sensíveis (pessoas, conflitos, dados) devem ser tratados pelos canais formais / RH / e-mail.",
            "Fora do horário de expediente, não deve esperar-se resposta imediata (ver secção 11).",
            "Antes de criar um novo ficheiro, verificar se já existe na pasta do canal.",
            "Decisões tomadas devem ser registadas de forma sucinta no canal.",
        ],
    )

    # 11
    add_heading(doc, "11. Notificações, telemóvel e disponibilidade")
    add_rich_paragraph(
        doc,
        "Uma preocupação frequente entre utilizadores habituados ao e-mail é o risco "
        "de interrupções constantes. Esse risco resulta sobretudo de notificações "
        "não configuradas. Recomenda-se o seguinte:",
    )
    add_bullets(
        doc,
        [
            "Definições → Notificações: manter alertas para **menções** e para o canal Avisos da respetiva AO.",
            "Silenciar o Team institucional se apenas forem necessários avisos pontuais, "
            "ou restringir as notificações a esse canal.",
            "No telemóvel, desativar a pré-visualização de conteúdo em locais públicos; utilizar autenticação.",
            "Utilizar o estado (Disponível, Ocupado, Não incomodar) de forma adequada à atividade em curso.",
            "A AO pode definir o seguinte compromisso: resposta no Teams em horário de expediente, "
            "sem obrigação fora desse período, salvo urgência devidamente comunicada pelo diretor.",
        ],
    )

    # 12
    add_heading(doc, "12. Plano de implementação (90 dias)")
    add_table(
        doc,
        ["Período", "Atividades", "Responsáveis"],
        [
            ["Semana 0", "Confirmação de licenças; criação dos Teams por AO; disponibilização deste manual no canal Avisos.", "IT e diretores"],
            ["Semana 1", "Sessão formativa de 20 minutos por AO. Comunicado de arranque. Apresentação do endereço de e-mail do canal Avisos. Projeto-piloto, se aplicável.", "Diretor e embaixador"],
            ["Semanas 2–4", "Comunicação operacional no Teams. Remissão educativa de e-mails operacionais para o canal.", "Toda a AO"],
            ["Dia 30", "Revisão: canais em utilização; ruído de notificações; utilização inadequada do Team institucional.", "Diretor"],
            ["Dia 90", "Avaliação: redução do e-mail operacional; reuniões da AO no Teams; integração de novos colaboradores.", "Direção e diretores de AO"],
        ],
    )
    add_caption(doc, "Quadro 5 — Calendário de implementação (adaptável à realidade da AIP).")
    add_rich_paragraph(
        doc,
        "Indicadores sugeridos (sem caráter de monitorização individual): número de canais "
        "com atividade útil; realização de reuniões semanais da AO no Teams; "
        "feedback qualitativo aos 30 e 90 dias. Evitar rankings de volume de mensagens.",
    )

    # 13
    add_heading(doc, "13. Perguntas frequentes")
    faqs = [
        (
            "O Team da minha AO não está visível. Como proceder?",
            "Solicitar a inclusão ao diretor ou ao IT. Evitar, entretanto, o recurso sistemático "
            "ao e-mail operacional, salvo urgência devidamente justificada.",
        ),
        (
            "É permitido continuar a utilizar o e-mail com os colegas?",
            "Sim, para assuntos formais, comunicação externa ou arquivo. "
            "Para a comunicação operacional interna, o padrão passa a ser o Teams.",
        ),
        (
            "O Team institucional será eliminado?",
            "Não necessariamente. Mantém-se para avisos transversais. "
            "A atividade de cada AO deixa de decorrer nesse espaço.",
        ),
        (
            "Qual a posição relativamente ao WhatsApp?",
            "Não constitui ferramenta oficial para comunicação operacional. "
            "O trabalho deve decorrer no Teams, com contas profissionais, ficheiros e histórico.",
        ),
        (
            "É obrigatório instalar o Teams no telemóvel?",
            "Não. É recomendável para colaboradores em deslocação. "
            "No posto de trabalho, a aplicação de ambiente de trabalho é suficiente.",
        ),
        (
            "O e-mail oferece maior valor probatório?",
            "As conversas e ficheiros do Teams também ficam registados na conta da organização. "
            "Quando a lei ou o processo exigirem documento formal, utiliza-se o e-mail ou o "
            "sistema oficial — não o e-mail para toda a comunicação «por precaução».",
        ),
        (
            "Posso criar um canal por iniciativa própria?",
            "Por princípio, deve solicitar-se ao diretor. "
            "O excesso de canais prejudica a utilização tanto quanto a sua escassez.",
        ),
        (
            "Como proceder face a um e-mail de cliente?",
            "Responder por e-mail. Caso seja necessário o envolvimento da AO, utilizar "
            "**Partilhar no Teams** para o canal adequado (secção 7.2), sem incluir o cliente no Team "
            "e sem reencaminhar cadeias com informação excessiva.",
        ),
        (
            "Como enviar o mesmo comunicado por e-mail e para o canal?",
            "No Outlook, incluir em Cc o endereço de e-mail do canal (secção 7.1). "
            "Se o e-mail já foi enviado, utilizar Partilhar no Teams. "
            "Não duplicar a comunicação operacional corrente — apenas comunicados relevantes.",
        ),
        (
            "Uma resposta no Outlook a um e-mail também enviado ao canal aparece no Teams?",
            "Em regra, não de forma fiável. As respostas no correio permanecem no fio de e-mail; "
            "as do canal permanecem no Teams. Deve definir-se à partida o canal de continuação "
            "(preferencialmente o canal do Teams).",
        ),
    ]
    for q, a in faqs:
        add_heading(doc, q, 2)
        add_rich_paragraph(doc, a)

    # 14
    add_heading(doc, "14. Glossário")
    add_table(
        doc,
        ["Termo", "Significado no âmbito deste modelo"],
        [
            ["Team (equipa)", "Espaço de trabalho. Na AIP, em regra: um por AO."],
            ["Canal", "Espaço temático dentro do Team (Geral, Avisos, projeto, entre outros)."],
            ["Conversa / chat", "Mensagem direta a uma pessoa ou a um grupo restrito."],
            ["@menção", "Notificação dirigida a uma pessoa ou ao canal."],
            ["Thread / resposta", "Sequência de mensagens sobre o mesmo assunto."],
            ["SharePoint", "Repositório onde residem os ficheiros do Team."],
            ["AO", "Área operacional — unidade de organização deste modelo."],
            ["Embaixador", "Colaborador que apoia a adoção na AO."],
            ["Endereço de e-mail do canal", "Endereço que permite publicar no canal mediante envio de e-mail."],
            ["Partilhar no Teams", "Funcionalidade do Outlook que remete um e-mail existente para um canal ou conversa."],
        ],
    )
    add_caption(doc, "Quadro 6 — Vocabulário de referência.")
    add_callout(
        doc,
        "**Próximo passo.** Compete ao diretor da AO agendar a sessão formativa de 20 minutos, "
        "publicar o comunicado de arranque e fixar no canal o critério "
        "«operacional → Teams; formal/externo → e-mail». "
        "A adoção consolida-se pela consistência de utilização, não por um anúncio isolado.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Documento de uso interno. Redigido em português de Portugal. "
        "Ferramentas: Microsoft Teams e Outlook / Microsoft 365."
    )
    _set_run_font(run, size=9, italic=True, color=MUTED)

    doc.save(str(path))
    return path


if __name__ == "__main__":
    print(build_docx())
